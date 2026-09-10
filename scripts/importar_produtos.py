"""Importa/atualiza produtos (schema do site flora-5754a) de um CSV ou XLSX.

Colunas esperadas (cabecalho):
    codigo_barras, nome, categoria, preco_varejo, preco_atacado,
    estoque_varejo, estoque_atacado, peso, descricao, ativo

- Upsert por `codigoBarras` (EAN) — a chave natural do produto no site.
- `codigo_barras` OBRIGATORIO: 8 a 14 digitos. Linha sem EAN valido e ignorada.
- `categoria` = slug de UMA opcao da camada principal de filtros (ex.: "perfumes",
  "iphones"). Gravado em `categoria` (legado) e em `filtros[<camada principal>]`.
  O slug da camada principal e lido da colecao `camadas` (menor `ordem`).
- `ativo`: "sim"/"nao"/"1"/"0"/vazio (default: sim).

Uso:
    python scripts/importar_produtos.py --arquivo scripts/exemplo_produtos.csv
"""
import argparse
import csv
import re

from firebase_admin import firestore

from _firebase import init

EAN_RE = re.compile(r"^\d{8,14}$")


def linhas(caminho):
    if caminho.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(caminho, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it)]
        for row in it:
            yield {head[i]: row[i] for i in range(len(head))}
    else:
        with open(caminho, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                yield row


def num(v, default=0.0):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return default


def sim_nao(v, default=True):
    s = str(v or "").strip().lower()
    if s in ("", "sim", "s", "1", "true", "ativo"):
        return default if s == "" else True
    if s in ("nao", "não", "n", "0", "false", "inativo"):
        return False
    return default


def slug_camada_principal(db):
    docs = list(db.collection("camadas").order_by("ordem").limit(1).stream())
    return (docs[0].to_dict().get("slug") if docs else None) or None


def main():
    ap = argparse.ArgumentParser(description="Importa produtos (schema do site).")
    ap.add_argument("--arquivo", required=True)
    ap.add_argument("--cred", default=None)
    args = ap.parse_args()

    db = init(args.cred)
    col = db.collection("produtos")
    camada = slug_camada_principal(db)
    if not camada:
        print("Aviso: nenhuma camada de filtro cadastrada — `filtros` ficara vazio "
              "e os produtos so terao o campo legado `categoria`.")

    novos = atualizados = ignorados = 0
    batch = db.batch()
    pendentes = 0

    for row in linhas(args.arquivo):
        ean = re.sub(r"\D", "", str(row.get("codigo_barras") or ""))
        if not EAN_RE.match(ean):
            print(f"  ignorado (EAN invalido): {row.get('codigo_barras')!r} / {row.get('nome')!r}")
            ignorados += 1
            continue

        categoria = str(row.get("categoria") or "").strip().lower()
        preco_atacado = num(row.get("preco_atacado"))

        dados = {
            "codigoBarras": ean,
            "nome": str(row.get("nome") or "").strip(),
            "descricao": str(row.get("descricao") or "").strip(),
            "categoria": categoria,
            "filtros": {camada: [categoria]} if (camada and categoria) else {},
            "precoVarejo": num(row.get("preco_varejo")),
            "precoAtacado": preco_atacado if preco_atacado > 0 else None,
            "estoqueVarejo": int(num(row.get("estoque_varejo"))),
            "estoqueAtacado": int(num(row.get("estoque_atacado"))),
            "estoque": None,
            "peso": int(num(row.get("peso"))),
            "ativo": sim_nao(row.get("ativo")),
            "atualizadoEm": firestore.SERVER_TIMESTAMP,
        }

        existente = next(iter(col.where("codigoBarras", "==", ean).limit(1).stream()), None)
        if existente:
            batch.set(existente.reference, dados, merge=True)
            atualizados += 1
        else:
            dados["criadoEm"] = firestore.SERVER_TIMESTAMP
            dados["destaque"] = False
            dados["freteDisponivel"] = True
            dados["imagemURL"] = ""
            dados["imagensExtras"] = []
            batch.set(col.document(), dados)
            novos += 1

        pendentes += 1
        if pendentes >= 400:
            batch.commit()
            batch = db.batch()
            pendentes = 0

    if pendentes:
        batch.commit()

    print(f"Novos: {novos} | Atualizados: {atualizados} | Linhas ignoradas: {ignorados}")


if __name__ == "__main__":
    main()
